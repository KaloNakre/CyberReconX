#!/usr/bin/env python3
"""
Generate a minimal PPTX and PDF presentation for the repository.

Slides structure (per code file):
 - Slide with code (monospace)
 - Slide with short explanation of how the code runs

This script produces `presentations/presentation.pptx` and
`presentations/presentation.pdf` under the repository root.
"""
import os
import textwrap

FILES = [
    "main.py",
    "install.sh",
    "install.ps1",
    "requirements.txt",
    "java/CyberReconX.java",
]


def read_file(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    except Exception:
        return ""


def make_pptx(out_path):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception:
        print("python-pptx not available; skipping PPTX generation")
        return False
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "CyberReconX — Code walkthrough"
    slide.placeholders[1].text = "Contains code snippets and short run explanations."

    for file in FILES:
        content = read_file(file)
        if not content:
            continue

        # Code slide
        slide = prs.slides.add_slide(blank_slide_layout)
        left = Inches(0.5)
        top = Inches(0.5)
        width = Inches(9)
        height = Inches(6.5)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.font.name = 'Courier New'
        p.font.size = Pt(12)
        # limit long files to first N lines for readability
        lines = content.splitlines()
        MAX_LINES = 60
        excerpt = "\n".join(lines[:MAX_LINES])
        if len(lines) > MAX_LINES:
            excerpt += "\n\n... (truncated)"
        tf.text = f"File: {file}\n\n{excerpt}"

        # Explanation slide
        expl = summarize_explanation(file, content)
        slide = prs.slides.add_slide(blank_slide_layout)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.font.size = Pt(14)
        p.text = f"How {file} runs:\n\n{expl}"

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return True


def make_pdf(out_path):
    # create a simple PDF with same content using reportlab
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import inch
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    c = canvas.Canvas(out_path, pagesize=letter)
    width, height = letter

    def write_page(title, text):
        c.setFont("Helvetica-Bold", 16)
        c.drawString(0.5 * inch, height - 1 * inch, title)
        c.setFont("Courier", 9)
        y = height - 1.5 * inch
        max_lines = 60
        for i, line in enumerate(text.splitlines()[:max_lines]):
            c.drawString(0.5 * inch, y - i * 12, line[:120])
        c.showPage()

    for file in FILES:
        content = read_file(file)
        if not content:
            continue
        excerpt = "\n".join(content.splitlines()[:60])
        write_page(f"File: {file}", excerpt)
        expl = summarize_explanation(file, content)
        write_page(f"How {file} runs", expl)

    c.save()


def summarize_explanation(file, content):
    # Minimal, deterministic explanations for the known files
    file = file.lower()
    if file.endswith("main.py"):
        return (
            "Interactive Python CLI: shows a banner, offers local or custom network recon, "
            "uses system tools (ifconfig/ipcalc) to detect local network, then runs nmap scans; "
            "OSINT uses whois and whatweb. Outputs are printed to stdout."
        )
    if file.endswith("install.sh"):
        return (
            "Bash installer for Debian-based systems: updates apt and installs python3, "
            "virtualenv dependencies and network tools (nmap, whois, whatweb, ipcalc, net-tools)."
        )
    if file.endswith("install.ps1"):
        return (
            "PowerShell installer: checks Python availability, creates a virtualenv, activates it, "
            "and installs Python dependencies from requirements.txt."
        )
    if file.endswith("requirements.txt"):
        return "Lists Python package dependencies for the project (may be empty if none required)."
    if file.endswith("cyberreconx.java"):
        return (
            "A minimal Java starter 'CyberReconX' that prints a small banner; added as a placeholder "
            "per request and does not interact with the Python code.")
    return "Short explanation not available."


def main():
    out_pptx = "presentations/presentation.pptx"
    out_pdf = "presentations/presentation.pdf"
    make_pptx(out_pptx)
    make_pdf(out_pdf)
    print("Generated:", out_pptx, out_pdf)


if __name__ == "__main__":
    main()
